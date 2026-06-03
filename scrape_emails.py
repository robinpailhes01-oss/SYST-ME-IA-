#!/usr/bin/env python3
"""Scrape company websites from leads.xlsx and add an EMAIL column."""
import re, sys, html, time, json
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urljoin, urlparse
import requests
from bs4 import BeautifulSoup
import openpyxl

INPUT = "leads.xlsx"
OUTPUT = "leads_with_emails.xlsx"
PROGRESS = "scrape_progress.json"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/120.0 Safari/537.36",
    "Accept-Language": "fr-FR,fr;q=0.9,en;q=0.8",
}

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")

# pages likely to contain an email
CANDIDATE_PATHS = [
    "", "contact", "contact/", "contactez-nous", "contactez-nous/",
    "nous-contacter", "mentions-legales", "mentions-legales/",
    "mentions-legales.html", "a-propos", "qui-sommes-nous", "legal",
    "contact.html", "contact.php",
]

BAD_SUBSTR = ("sentry", "wixpress", "wix.com", "godaddy", "example.",
              "domain.com", "yourdomain", "email@", "@2x", "@sentry",
              ".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".bmp",
              "@adobe", "@schema", "u003", "%", "sentry.io", "core.js",
              "local.fr", "etre-visible", ".local", "wordpress", "wix.png",
              "@x.com", "@2.x", "no-reply@", "noreply@", "@email.com")


def clean_emails(found, domain):
    out = []
    for e in found:
        e = html.unescape(e).strip().strip(".,;:'\"()<>").lower()
        if not e or "@" not in e:
            continue
        low = e.lower()
        if any(b in low for b in BAD_SUBSTR):
            continue
        if low.split("@")[-1].count(".") == 0:
            continue
        # skip emails that are clearly hashed/asset names
        local = low.split("@")[0]
        if len(local) > 40 or len(low) > 60:
            continue
        out.append(e)
    # dedupe preserving order
    seen, uniq = set(), []
    for e in out:
        if e not in seen:
            seen.add(e); uniq.append(e)
    return uniq


def decode_cfemail(soup):
    """Decode Cloudflare-obfuscated emails (data-cfemail attributes)."""
    res = []
    for tag in soup.select("[data-cfemail]"):
        enc = tag.get("data-cfemail")
        try:
            r = int(enc[:2], 16)
            dec = "".join(chr(int(enc[i:i+2], 16) ^ r) for i in range(2, len(enc), 2))
            res.append(dec)
        except Exception:
            pass
    return res


def fetch(url, session):
    try:
        r = session.get(url, headers=HEADERS, timeout=12, allow_redirects=True)
        ct = r.headers.get("Content-Type", "")
        if r.status_code == 200 and ("html" in ct or "text" in ct or ct == ""):
            return r.text
    except Exception:
        pass
    return None


def extract_from_html(text, domain):
    emails = []
    soup = BeautifulSoup(text, "html.parser")
    # mailto links
    for a in soup.select('a[href^="mailto:"]'):
        href = a.get("href", "")
        addr = href[len("mailto:"):].split("?")[0]
        emails.extend(EMAIL_RE.findall(html.unescape(addr)))
    # cloudflare
    emails.extend(decode_cfemail(soup))
    # raw text regex
    emails.extend(EMAIL_RE.findall(text))
    return clean_emails(emails, domain)


def pick_best(emails, domain):
    if not emails:
        return ""
    dom = domain.lower().replace("www.", "")
    # prefer same-domain
    same = [e for e in emails if e.split("@")[-1].replace("www.", "") == dom]
    pool = same if same else emails
    pref = ("contact@", "info@", "bonjour@", "hello@", "accueil@", "commercial@")
    for p in pref:
        for e in pool:
            if e.startswith(p):
                return e
    return pool[0]


def scrape_site(url):
    """Return (best_email, all_emails_joined)."""
    if not url or not isinstance(url, str):
        return "", ""
    url = url.strip()
    if not url.startswith("http"):
        url = "http://" + url
    parsed = urlparse(url)
    base = f"{parsed.scheme}://{parsed.netloc}"
    domain = parsed.netloc
    session = requests.Session()

    # Build ordered candidate list: the given URL first, then base + paths
    candidates = [url]
    for p in CANDIDATE_PATHS:
        c = urljoin(base + "/", p)
        if c not in candidates:
            candidates.append(c)

    all_emails = []
    seen_pages = set()
    for c in candidates:
        if c in seen_pages:
            continue
        seen_pages.add(c)
        text = fetch(c, session)
        if text:
            found = extract_from_html(text, domain)
            for e in found:
                if e not in all_emails:
                    all_emails.append(e)
            # discover contact/legal links from homepage if still nothing
            if not all_emails and c in (url, base + "/", base):
                soup = BeautifulSoup(text, "html.parser")
                for a in soup.find_all("a", href=True):
                    h = a["href"].lower()
                    if any(k in h for k in ("contact", "mention", "legal", "propos")):
                        link = urljoin(c, a["href"])
                        if urlparse(link).netloc == domain and link not in seen_pages:
                            candidates.append(link)
        # stop early once we have a good same-domain email
        if all_emails and any(e.split("@")[-1].replace("www.", "") ==
                              domain.lower().replace("www.", "") for e in all_emails):
            break
        if len(seen_pages) >= 8:
            break

    best = pick_best(all_emails, domain)
    return best, "; ".join(all_emails)


def main():
    wb = openpyxl.load_workbook(INPUT)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    # header
    ws.cell(row=1, column=4, value="MAIL")
    ws.cell(row=1, column=5, value="MAILS_TROUVES")

    tasks = []  # (excel_row, url)
    for idx, r in enumerate(rows):
        excel_row = idx + 1
        if idx == 0:
            continue  # header row
        url = r[1] if len(r) > 1 else None
        if url and isinstance(url, str) and url.strip():
            tasks.append((excel_row, url.strip()))

    total = len(tasks)
    print(f"Scraping {total} sites...", flush=True)
    done = 0
    results = {}
    with ThreadPoolExecutor(max_workers=24) as ex:
        fut = {ex.submit(scrape_site, url): (er, url) for er, url in tasks}
        for f in as_completed(fut):
            er, url = fut[f]
            try:
                best, allm = f.result()
            except Exception as e:
                best, allm = "", ""
            results[er] = (best, allm)
            done += 1
            if done % 25 == 0 or done == total:
                found = sum(1 for v in results.values() if v[0])
                print(f"  {done}/{total} done, {found} emails found", flush=True)
                json.dump({"done": done, "total": total, "found": found},
                          open(PROGRESS, "w"))

    for er, (best, allm) in results.items():
        ws.cell(row=er, column=4, value=best)
        ws.cell(row=er, column=5, value=allm)

    wb.save(OUTPUT)
    found = sum(1 for v in results.values() if v[0])
    print(f"DONE. {found}/{total} emails found. Saved {OUTPUT}", flush=True)
    json.dump({"done": total, "total": total, "found": found, "complete": True},
              open(PROGRESS, "w"))


if __name__ == "__main__":
    main()
